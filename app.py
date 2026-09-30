import streamlit as st
import pandas as pd
import openpyxl

# Page Configuration
st.set_page_config(page_title="RHS Portal - Hardware & Paints", page_icon="🏢", layout="wide", initial_sidebar_state="collapsed")

# ----------------- FLIPKART STYLE 3D CUSTOM CSS -----------------
st.markdown("""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;600;700;800&display=swap');
    html, body, [class*="css"] { font-family: 'Inter', sans-serif; }
    .stApp { background-color: #F1F3F6; }
    .fk-navbar {
        background: linear-gradient(135deg, #2874F0 0%, #1A56B8 100%);
        padding: 14px 22px; border-radius: 12px; color: white;
        display: flex; align-items: center; justify-content: space-between;
        box-shadow: 0 8px 20px rgba(40, 116, 240, 0.25); margin-bottom: 20px;
    }
    .fk-navbar h2 { margin: 0; font-size: 20px; font-weight: 800; letter-spacing: 0.5px; }
    .badge-admin { background: #FEF3C7; color: #92400E; padding: 4px 10px; border-radius: 20px; font-weight: 700; font-size: 12px; }
    .badge-public { background: #ECFDF5; color: #065F46; padding: 4px 10px; border-radius: 20px; font-weight: 700; font-size: 12px; }
    div[data-baseweb="input"] { border-radius: 10px !important; border: 2px solid #E2E8F0 !important; }
</style>
""", unsafe_allow_html=True)

# ----------------- ADMIN CREDENTIALS & SESSION STATE -----------------
ADMIN_USERS = {"admin": "rhspaint2026", "owner": "radha@sagar"}

if "is_admin" not in st.session_state: st.session_state.is_admin = False
if "gst_rate" not in st.session_state: st.session_state.gst_rate = 14.0
if "scheme_val" not in st.session_state: st.session_state.scheme_val = 10.0
if "paint_items" not in st.session_state: st.session_state.paint_items = []
if "hw_items" not in st.session_state: st.session_state.hw_items = []

# ----------------- SIDEBAR (ADMIN CONTROLS) -----------------
with st.sidebar:
    st.markdown("### 🔐 Admin Controls")
    if not st.session_state.is_admin:
        u_name = st.text_input("Username")
        u_pass = st.text_input("Password", type="password")
        if st.button("Login", use_container_width=True):
            if u_name in ADMIN_USERS and ADMIN_USERS[u_name] == u_pass:
                st.session_state.is_admin = True
                st.rerun()
            else:
                st.error("Invalid Credentials")
    else:
        st.markdown('<span class="badge-admin">Admin Active</span>', unsafe_allow_html=True)
        if st.button("Logout", use_container_width=True):
            st.session_state.is_admin = False
            st.rerun()
        
        st.markdown("---")
        st.subheader("🎨 Paint Rates Config")
        n_gst = st.number_input("GST Rate (%)", value=st.session_state.gst_rate)
        n_scheme = st.number_input("Scheme Per Litre (₹)", value=st.session_state.scheme_val)
        if st.button("Update Paint Rates"):
            st.session_state.gst_rate, st.session_state.scheme_val = n_gst, n_scheme
            st.success("Updated!")
            st.rerun()

        st.markdown("---")
        st.subheader("📁 Upload Any Price List")
        uploaded_file = st.file_uploader("Upload Excel (.xlsx)", type=["xlsx"])
        
        if uploaded_file and st.button("Process & Import"):
            try:
                wb = openpyxl.load_workbook(uploaded_file, data_only=True)
                ws = wb.active
                
                # Check Header Type (Row 1)
                headers = [str(ws.cell(row=1, column=c).value or "").strip().lower() for c in range(1, 7)]
                
                # --- HARDWARE UPLOAD LOGIC ---
                if "whole sale" in headers or "retail price" in headers:
                    hw_map = {(i["company"], i["name"], i["size"]): i for i in st.session_state.hw_items}
                    added, updated = 0, 0
                    
                    for r in range(2, ws.max_row + 1):
                        c1 = ws.cell(row=r, column=1).value # Company
                        c2 = ws.cell(row=r, column=2).value # Item Name
                        c3 = ws.cell(row=r, column=3).value # Pack
                        c4 = ws.cell(row=r, column=4).value # Size
                        c5 = ws.cell(row=r, column=5).value # Wholesale
                        c6 = ws.cell(row=r, column=6).value # Retail
                        
                        if c2 and c5 and c6: # Requires Name and Prices
                            comp = str(c1).strip().upper() if c1 else "-"
                            name = str(c2).strip().upper()
                            pack = str(c3).strip() if c3 else "-"
                            size = str(c4).strip() if c4 else "-"
                            
                            try:
                                w_price, r_price = float(c5), float(c6)
                                key = (comp, name, size)
                                if key in hw_map:
                                    hw_map[key].update({"pack": pack, "wholesale": w_price, "retail": r_price})
                                    updated += 1
                                else:
                                    hw_map[key] = {"company": comp, "name": name, "pack": pack, "size": size, "wholesale": w_price, "retail": r_price}
                                    added += 1
                            except ValueError: continue
                            
                    st.session_state.hw_items = list(hw_map.values())
                    st.success(f"Hardware List: {added} Added, {updated} Updated!")
                
                # --- PAINTS UPLOAD LOGIC ---
                else:
                    paint_map = {(i["name"], i["pack"]): i for i in st.session_state.paint_items}
                    start_row = 2
                    for r in range(1, 15):
                        r_vals = [str(ws.cell(row=r, column=c).value or "").strip().upper() for c in range(1, 5)]
                        if any("NAME" in x for x in r_vals): start_row = r + 1; break
                        
                    added, updated = 0, 0
                    for r in range(start_row, ws.max_row + 1):
                        b, c, d = ws.cell(row=r, column=2).value, ws.cell(row=r, column=3).value, ws.cell(row=r, column=4).value
                        if b and c and d:
                            n_str = str(b).strip().upper()
                            try:
                                p_num, r_num = float(str(c).strip()), float(str(d).strip())
                                key = (n_str, p_num)
                                if key in paint_map:
                                    paint_map[key]["rate"] = r_num; updated += 1
                                else:
                                    paint_map[key] = {"name": n_str, "pack": p_num, "rate": r_num}; added += 1
                            except ValueError: continue
                            
                    st.session_state.paint_items = list(paint_map.values())
                    st.success(f"Paint List: {added} Added, {updated} Updated!")
            except Exception as e:
                st.error(f"Error: {e}")

# ----------------- MAIN UI -----------------
role = '<span class="badge-admin">ADMINISTRATOR</span>' if st.session_state.is_admin else '<span class="badge-public">COUNTER USER</span>'
st.markdown(f"""
<div class="fk-navbar">
    <div><h2>🏢 RHS PORTAL - HARDWARE & PAINTS</h2><small style="opacity: 0.9;">Universal Rate Search Engine</small></div>
    <div>{role}</div>
</div>
""", unsafe_allow_html=True)

search_query = st.text_input("🔍 Quick Global Search", placeholder="Type item name, company, or colour...").strip().upper()

tab1, tab2 = st.tabs(["🛠️ HARDWARE & GENERAL", "🎨 PAINTS & SCHEMES"])

# --- TAB 1: HARDWARE ---
with tab1:
    if st.session_state.hw_items:
        df_hw = pd.DataFrame(st.session_state.hw_items)
        df_hw.rename(columns={"company": "COMPANY", "name": "ITEM NAME", "pack": "PACK", "size": "SIZE", "wholesale": "WHOLESALE (₹)", "retail": "RETAIL (₹)"}, inplace=True)
        
        if search_query:
            df_hw = df_hw[df_hw["ITEM NAME"].str.contains(search_query, na=False) | df_hw["COMPANY"].str.contains(search_query, na=False)]
        
        if not st.session_state.is_admin:
            df_hw = df_hw.drop(columns=["WHOLESALE (₹)"]) # Hide wholesale from counter staff
            
        st.dataframe(df_hw.style.set_properties(**{'background-color': '#FFFFFF', 'font-weight': '600'}), use_container_width=True, height=400)
    else:
        st.info("Hardware list khali hai. Admin panel se 6-column wali excel upload karein.")

# --- TAB 2: PAINTS ---
with tab2:
    if st.session_state.paint_items:
        calc_rows = []
        for item in st.session_state.paint_items:
            rate_gst = round(item["rate"] * (1 + (st.session_state.gst_rate / 100)), 2)
            final_rate = round(rate_gst - (item["pack"] * st.session_state.scheme_val), 2)
            row = {"COLOUR NAME": item["name"], "PACK (L/Kg)": item["pack"], "FINAL RATE (₹)": final_rate}
            if st.session_state.is_admin:
                row["DEALER (₹)"] = item["rate"]
                row["WITH GST (₹)"] = rate_gst
            calc_rows.append(row)
            
        df_pt = pd.DataFrame(calc_rows)
        if search_query: df_pt = df_pt[df_pt["COLOUR NAME"].str.contains(search_query, na=False)]
        
        col1, col2 = st.columns(2)
        with col1:
            st.markdown("##### 📋 Item Breakdown")
            st.dataframe(df_pt.style.set_properties(**{'background-color': '#FFFFFF', 'color': '#000000'}), use_container_width=True, height=350)
        with col2:
            st.markdown("##### 🏪 Print Matrix Grid")
            if not df_pt.empty:
                pivot = df_pt.pivot_table(index="COLOUR NAME", columns="PACK (L/Kg)", values="FINAL RATE (₹)", aggfunc="first").round(2).fillna("-")
                st.dataframe(pivot.style.format(precision=2, na_rep="-").set_properties(**{'background-color': '#FFFDF7', 'color': '#000000', 'font-weight': '600'}), use_container_width=True)
    else:
        st.info("Paint list khali hai. Admin panel se paint wali excel upload karein.")
